"""生成 Quill 项目介绍 PPT（5 页，16:9）。

设计口径：
- 配色取「墨蓝 · 羊皮纸」品牌方案（quill-source/.workbuddy/memory/2026-09-19.md）
- 图标取 quill-source/png/Quill_1024.png（02a 详细版，小尺寸另有 02c 简化版）
- 文案来源：README.md / CHANGELOG.md / pyproject.toml，不臆造

用法：python deck/build_deck.py
输出：deck/quill-intro.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
ICON = ROOT / "quill-source" / "png" / "Quill_1024.png"
OUT = Path(__file__).resolve().parent / "quill-intro.pptx"

# ---- 品牌色（墨蓝 · 羊皮纸） ----
INK = RGBColor(0x1B, 0x3A, 0x5C)    # 主色 Ink Navy 墨蓝
DEEP = RGBColor(0x0F, 0x22, 0x33)   # 深色 Deep Ink 深墨
BLUE = RGBColor(0x3D, 0x6E, 0x9E)   # 辅色 Quill Blue 羽蓝
GOLD = RGBColor(0xC8, 0x93, 0x3F)   # 强调 Amber Gold 琥珀金
PARCH = RGBColor(0xF5, 0xF0, 0xE6)  # 浅底 Parchment 羊皮纸
SAND = RGBColor(0xE3, 0xD9, 0xC6)   # 中性 Warm Sand 暖沙
GRAPH = RGBColor(0x5A, 0x64, 0x72)  # 正文 Graphite 石墨灰
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xD4, 0xE0)   # 深底上的次级文字（墨蓝提亮）

CN = "PingFang SC"
EN = "Helvetica Neue"
MONO = "SF Mono"

SW, SH = Inches(13.333), Inches(7.5)
M = Inches(0.9)
CW = Inches(13.333) - Inches(0.9) * 2


def style_run(run, size=14, bold=False, color=GRAPH, latin=EN, ea=CN):
    """设置字号/颜色，并同时挂上拉丁字体与中日韩字体。"""
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", ea)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, first=False, space_after=6, space_before=0, line_spacing=1.25):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line_spacing
    return p


def add_runs(p, parts):
    """parts: [(text, size, bold, color, latin)]"""
    for text, size, bold, color, latin in parts:
        r = p.add_run()
        r.text = text
        style_run(r, size=size, bold=bold, color=color, latin=latin)
    return p


def rect(slide, x, y, w, h, fill, radius=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    if radius is not None:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp


def set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def content_header(slide, title, kicker=None):
    """内容页统一页眉：金色短竖线 + 标题 + 暖沙分隔线。"""
    rect(slide, M, Inches(0.55), Inches(0.055), Inches(0.42), GOLD, radius=0.5)
    tf = textbox(slide, M + Inches(0.22), Inches(0.5), CW - Inches(0.22), Inches(0.55))
    p = para(tf, first=True, space_after=0)
    add_runs(p, [(title, 25, True, INK, EN)])
    if kicker:
        tf2 = textbox(slide, M, Inches(1.06), CW, Inches(0.3))
        p2 = para(tf2, first=True, space_after=0)
        add_runs(p2, [(kicker, 12.5, False, GRAPH, EN)])
    rect(slide, M, Inches(1.42), CW, Inches(0.012), SAND)


def footer(slide, page, note=None):
    tf = textbox(slide, M, Inches(6.98), CW - Inches(0.6), Inches(0.3))
    p = para(tf, first=True, space_after=0)
    add_runs(p, [(note or "", 9.5, False, GRAPH, EN)])
    tf2 = textbox(slide, Inches(13.333) - M - Inches(0.6), Inches(6.98), Inches(0.6), Inches(0.3))
    p2 = para(tf2, first=True, space_after=0)
    p2.alignment = PP_ALIGN.RIGHT
    add_runs(p2, [(str(page), 9.5, True, INK, EN)])


def card(slide, x, y, w, h, title, lines, accent=GOLD, fill=WHITE, tsize=15, bsize=11.5):
    rect(slide, x, y, w, h, fill, radius=0.08)
    rect(slide, x, y + Inches(0.18), Inches(0.05), h - Inches(0.36), accent, radius=0.5)
    tf = textbox(slide, x + Inches(0.24), y + Inches(0.16), w - Inches(0.45), h - Inches(0.3))
    p = para(tf, first=True, space_after=7, line_spacing=1.15)
    add_runs(p, [(title, tsize, True, INK, EN)])
    for ln in lines:
        pl = para(tf, space_after=5, line_spacing=1.3)
        add_runs(pl, [(ln, bsize, False, GRAPH, EN)])


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    blank = prs.slide_layouts[6]

    # ================= 1 封面 =================
    s = prs.slides.add_slide(blank)
    set_bg(s, DEEP)
    rect(s, Inches(0), Inches(0), Inches(0.14), SH, GOLD, radius=0)  # 左侧金边
    s.shapes.add_picture(str(ICON), Inches(1.15), Inches(2.05), height=Inches(2.6))
    # 图标自带墨蓝圆角底，压在深墨背景上边界偏糊 —— 描一圈金边把轮廓提出来
    ring = rect(s, Inches(1.1), Inches(2.0), Inches(2.7), Inches(2.7), None, radius=0.21, line=GOLD)
    ring.line.width = Pt(1.25)

    tf = textbox(s, Inches(4.5), Inches(2.15), Inches(7.6), Inches(1.8))
    p = para(tf, first=True, space_after=4)
    add_runs(p, [("Quill", 54, True, WHITE, EN)])
    p = para(tf, space_after=0)
    add_runs(p, [("本地优先的可配置 Agent 应用", 20, False, GOLD, EN)])

    rect(s, Inches(4.5), Inches(4.35), Inches(1.1), Inches(0.03), GOLD)

    tf = textbox(s, Inches(4.5), Inches(4.7), Inches(7.9), Inches(1.2))
    p = para(tf, first=True, space_after=6)
    add_runs(p, [("业务层与界面完全分离：", 13.5, False, SOFT, EN),
                 ("提示词、工具、技能、记忆都可配可改", 13.5, True, WHITE, EN)])
    p = para(tf, space_after=0)
    add_runs(p, [("同一份纯 Python 逻辑，Vue 前端、命令行、脚本都能调", 13.5, False, SOFT, EN)])

    tf = textbox(s, Inches(4.5), Inches(6.35), Inches(7.9), Inches(0.4))
    p = para(tf, first=True, space_after=0)
    add_runs(p, [("v0.1.8  ·  2026-09  ·  MIT License", 11.5, False, SOFT, MONO)])

    # ================= 2 定位与架构 =================
    s = prs.slides.add_slide(blank)
    set_bg(s, PARCH)
    content_header(s, "可本地运行、可二次开发的 Agent 应用",
                   "三层各管一件事，界面换了业务层不用动")

    boxes = [
        (Inches(0.9), "web/", "Vue 3 + Element Plus", "界面层：对话、模式、提示词、技能页"),
        (Inches(5.35), "server/", "FastAPI", "接口层：HTTP / SSE 流式转发"),
        (
            Inches(9.8),
            "src/quill_agent/",
            "纯 Python · 无界面依赖",
            "业务层：Agent 循环、工具、记忆、持久化",
        ),
    ]
    for x, name, tech, desc in boxes:
        rect(s, x, Inches(1.78), Inches(2.8), Inches(1.45), WHITE, radius=0.1)
        tf = textbox(s, x + Inches(0.2), Inches(1.9), Inches(2.4), Inches(1.25))
        p = para(tf, first=True, space_after=3, line_spacing=1.1)
        add_runs(p, [(name, 14.5, True, INK, MONO)])
        p = para(tf, space_after=3, line_spacing=1.1)
        add_runs(p, [(tech, 11, True, GOLD, EN)])
        p = para(tf, space_after=0, line_spacing=1.2)
        add_runs(p, [(desc, 10, False, GRAPH, EN)])

    for x, label in ((Inches(3.85), "HTTP / SSE"), (Inches(8.3), "直接调用")):
        tf = textbox(s, x, Inches(2.14), Inches(1.4), Inches(0.55))
        p = para(tf, first=True, space_after=0)
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, [("───▶", 13, True, BLUE, EN)])
        p = para(tf, space_after=0)
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, [(label, 9.5, False, GRAPH, EN)])

    cw = (CW - Inches(0.8)) / 3
    card(s, M, Inches(3.55), cw, Inches(2.75), "为什么分层",
         ["业务层不 import 任何界面框架，只认 Python。",
          "今天被 Vue 前端调用，明天被 CLI 或脚本复用 —— 换界面不用重写逻辑。"],
         accent=GOLD)
    card(s, M + cw + Inches(0.4), Inches(3.55), cw, Inches(2.75), "数据在本地",
         ["会话、记忆、模式、提示词全是本地 JSON / JSONL 文件。",
          "检索走关键词匹配、零 embedding 依赖，对话内容不出本机；换模型只改配置。"],
         accent=BLUE)
    card(s, M + (cw + Inches(0.4)) * 2, Inches(3.55), cw, Inches(2.75), "可二次开发",
         ["加一个工具 = 加一个 @registry.tool 装饰器。",
          "改提示词、加技能 = 编辑 prompt/ 与 skills/ 下的 md 文件，存盘即生效。"],
         accent=GOLD)
    footer(s, 2, "来源：README.md 第 1 节 / 第 2 节 / 第 4 节")

    # ================= 3 设计逻辑 =================
    s = prs.slides.add_slide(blank)
    set_bg(s, PARCH)
    content_header(s, "一个「模式」，决定这一轮 Agent 的全部行为",
                   "模式 Mode（data/modes.json）= 提示词组 + 工具组 + 技能组 + 记忆开关 + 偏好模型")

    rect(s, M, Inches(1.72), CW, Inches(0.62), INK, radius=0.12)
    tf = textbox(s, M + Inches(0.3), Inches(1.72), CW - Inches(0.6), Inches(0.62),
                 anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, first=True, space_after=0)
    add_runs(p, [("请求 → 流式解析 → 执行工具 → 再请求", 14, True, WHITE, MONO),
                 ("     四类资源各自落到 system 消息、tools 参数与 user 消息上",
                  11.5, False, SOFT, EN)])

    cw2 = (CW - Inches(0.4)) / 2
    ch2 = Inches(1.72)
    card(s, M, Inches(2.58), cw2, ch2, "提示词 · 每轮全量注入",
         ["六类按固定顺序拼进 system：身份 / 能力 / 工具策略 / 工作流程 / 输出规范 / 约束。",
          "正文是文件（prompt/<id>.md），组合是配置 —— 提示词组从每类里各挑一个。"],
         accent=GOLD, tsize=14, bsize=11)
    card(s, M + cw2 + Inches(0.4), Inches(2.58), cw2, ch2, "工具 · 按组下发",
         ["共 27 个：文件读写、命令执行、后台进程、规划、交互、技能、子代理……",
          "不在模式工具组里的工具不会被发给模型，模型自然不会请求调用。"],
         accent=BLUE, tsize=14, bsize=11)
    card(s, M, Inches(4.48), cw2, ch2, "技能 · 按需加载",
         ["平时只占一行「名字 + 适用场景」（约 30 字），任务匹配时模型自己 read_skill 读全文。",
          "20 个流程说明从 1.6 万字常驻开销，变成用到才付费。"],
         accent=BLUE, tsize=14, bsize=11)
    card(s, M + cw2 + Inches(0.4), Inches(4.48), cw2, ch2, "记忆 · 三个问题分开解",
         ["工作记忆 = 历史按窗口预算截断、工具调用可还原；程序性记忆 = 技能。",
          "长期记忆 = data/memory.json + remember 工具；截掉的原文可用 recall_history 捞回。"],
         accent=GOLD, tsize=14, bsize=11)
    footer(
        s,
        3,
        "来源：README.md 第 3 节（3.1 提示词 / 3.2 工具 / 3.3 技能 / 3.4 记忆 / 3.7 模式）",
    )

    # ================= 4 可控与安全 =================
    s = prs.slides.add_slide(blank)
    set_bg(s, PARCH)
    content_header(s, "边界说清楚：能碰什么、要不要问你",
                   "授权、审批、沙箱是三道互补的锁 —— 审批挡不住「不小心」，沙箱才能")

    cw = (CW - Inches(0.8)) / 3
    card(s, M, Inches(1.8), cw, Inches(2.9), "① 授权 · 应用层",
         ["工具组是白名单：没加进模式工具组的工具，根本不会下发给模型。",
          "一条命令要不要放行，取决于它所在的工具在不在组里。"],
         accent=GOLD, tsize=14.5)
    card(s, M + cw + Inches(0.4), Inches(1.8), cw, Inches(2.9), "② 审批 · 应用层",
         ["工具组里标了 confirm 的操作会弹给用户，阻塞着等答案。",
          "同意才执行；拒绝或问不到，就把结论当工具结果回填给模型继续跑。"],
         accent=BLUE, tsize=14.5)
    card(s, M + (cw + Inches(0.4)) * 2, Inches(1.8), cw, Inches(2.9), "③ 沙箱 · 内核层",
         ["SANDBOX_MODE 三档：off 全盘 / read-only 只读 / workspace-write 只给工作目录。",
          "Linux 走 bubblewrap，物理限制越界写入；网络默认关闭，要开得显式加参数。"],
         accent=GOLD, tsize=14.5)

    rect(s, M, Inches(5.05), CW, Inches(1.35), WHITE, radius=0.1)
    tf = textbox(s, M + Inches(0.3), Inches(5.2), CW - Inches(0.6), Inches(1.05))
    p = para(tf, first=True, space_after=5)
    add_runs(p, [("数据同样可控", 13.5, True, INK, EN)])
    p = para(tf, space_after=0, line_spacing=1.3)
    add_runs(p, [("会话、记忆、技能、提示词全部落在本地文件里：不引第三方检索服务，不上传对话内容；"
                  "会话作为产出物可导出带走，删一个工具/技能也不会让对话跑不起来。",
                  11.5, False, GRAPH, EN)])
    footer(s, 4, "来源：README.md 3.9（让 Agent 停下来等人）/ 3.10（沙箱）/ 5（持久层）")

    # ================= 5 快速开始 =================
    s = prs.slides.add_slide(blank)
    set_bg(s, PARCH)
    content_header(s, "怎么跑起来", "一条命令起步，五分钟后就能对话")

    rect(s, M, Inches(1.8), Inches(7.2), Inches(3.5), DEEP, radius=0.06)
    tf = textbox(s, M + Inches(0.32), Inches(2.02), Inches(6.6), Inches(3.1))
    code = [
        ("# 最省事：双击启动脚本（会装依赖、起服务、开浏览器）", SOFT, False),
        ("start.command", GOLD, True),
        ("            # macOS；Windows 用 start.bat，Linux 跑 ./start.sh", SOFT, False),
        ("", WHITE, False),
        ("# 或者自己掌舵", SOFT, False),
        ("make dev", GOLD, True),
        ("                 # 安装后端依赖（uv 自动建 .venv）", SOFT, False),
        ("make build && make api", GOLD, True),
        ("      # 构建前端并启动 → localhost:8000", SOFT, False),
        ("uv run quill serve --open-browser", GOLD, True),
        ("   # 等价写法：端口占用自动顺延", SOFT, False),
    ]
    for i, (txt, color, bold) in enumerate(code):
        p = para(tf, first=(i == 0), space_after=2, line_spacing=1.15)
        add_runs(p, [(txt, 12, bold, color, MONO)])

    rect(s, M + Inches(7.6), Inches(1.8), CW - Inches(7.6), Inches(1.62), WHITE, radius=0.1)
    tf = textbox(s, M + Inches(7.85), Inches(1.95), CW - Inches(8.1), Inches(1.35))
    p = para(tf, first=True, space_after=5)
    add_runs(p, [("环境要求", 13, True, INK, EN)])
    for line in ["Python ≥ 3.10（开发用 3.12）", "包管理 uv —— 仓库含 uv.lock，安装可复现",
                 "Node ≥ 20（只在构建前端时需要）"]:
        pl = para(tf, space_after=3, line_spacing=1.2)
        add_runs(pl, [("· " + line, 10.5, False, GRAPH, EN)])

    rect(s, M + Inches(7.6), Inches(3.62), CW - Inches(7.6), Inches(1.68), WHITE, radius=0.1)
    tf = textbox(s, M + Inches(7.85), Inches(3.77), CW - Inches(8.1), Inches(1.4))
    p = para(tf, first=True, space_after=5)
    add_runs(p, [("当前状态", 13, True, INK, EN)])
    for line in ["v0.1.8（2026-09-21）· MIT 许可", "macOS / Linux / Windows 均可运行",
                 "沙箱隔离目前只有 Linux / WSL2 有后端"]:
        pl = para(tf, space_after=3, line_spacing=1.2)
        add_runs(pl, [("· " + line, 10.5, False, GRAPH, EN)])

    tf = textbox(s, M, Inches(5.6), CW, Inches(0.9))
    p = para(tf, first=True, space_after=4, line_spacing=1.25)
    add_runs(p, [("接着可以试：", 12, True, INK, EN),
                 ("在「模式」页建一个自己的模式 —— 挑一组提示词、勾几个工具、挂一个技能，"
                  "然后回任务页选它开跑。", 12, False, GRAPH, EN)])
    footer(s, 5, "来源：README.md 第 1 节 / pyproject.toml / CHANGELOG.md 0.1.8 条目")

    prs.save(str(OUT))
    print(f"已生成：{OUT}，共 {len(prs.slides)} 页")


if __name__ == "__main__":
    build()
